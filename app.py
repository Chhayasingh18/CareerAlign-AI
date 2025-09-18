import gradio as gr
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
import requests
from collections import Counter
import re

SKILLS_DB = [
    # Core Programming & AI Skills
    "Python", "Java", "C++", "JavaScript", "SQL", "R", "MATLAB",
    "Machine Learning", "Deep Learning", "Data Analysis", "Pandas", "NumPy",
    "TensorFlow", "PyTorch", "scikit-learn", "Statistics", "Data Visualization",
    # Cloud & DevOps
    "AWS", "Azure", "Google Cloud", "Docker", "Kubernetes", "API Development",
    # Web/App Development
    "HTML", "CSS", "Sass", "React", "Angular", "Vue", "Next.js", "Django", "Flask",
    "Firebase", "WordPress",
    # Collaboration & Workflow
    "Git", "GitHub", "JIRA", "Slack", "Trello", "Notion", "Asana",
    # Productivity & Office
    "Excel", "Google Sheets", "PowerPoint", "Microsoft Word",
    # Content Creation Tools
    "Canva", "Figma", "Photoshop", "Illustrator", "Premiere Pro", "Final Cut Pro",
    "DaVinci Resolve", "Audacity", "OBS Studio", "Screen Recording", "Video Editing",
    "Image Editing", "Audio Editing", "Animation", "Motion Graphics",
    # AI & Generative Media
    "Prompt Engineering", "Natural Language Processing", "ChatGPT", "Llama",
    "Generative AI", "LangChain", "Stable Diffusion", "Midjourney", "D-ID",
    "Speech-to-Text", "Text-to-Speech", "Face Animation", "Voice Cloning", "OpenCV",
    "Text Summarization", "Topic Modeling", "Automated Captioning",
    # Content, Marketing, Publishing
    "Content Automation", "Content Repurposing", "Newsletter Automation",
    "Blog Writing", "SEO", "Digital Marketing", "Branding", "Copywriting",
    "Social Media Management", "YouTube Studio", "Podcast Editing", "Online Courses",
    "Analytics", "Google Analytics", "Monetization", "Affiliate Marketing",
    "Email Marketing", "E-commerce Platforms",
    # Other creative skills
    "Storytelling", "Scriptwriting", "Photo Retouching", "Video Thumbnails",
    "Livestreaming", "Online Collaboration", "Campaign Optimization"
]



RESOURCES = {
    # Core Programming & AI
    "Python": "https://www.learnpython.org/",
    "Java": "https://www.learnjavaonline.org/",
    "C++": "https://www.learncpp.com/",
    "JavaScript": "https://javascript.info/",
    "SQL": "https://www.w3schools.com/sql/",
    "R": "https://cran.r-project.org/manuals.html",
    "MATLAB": "https://matlabacademy.mathworks.com/",
    "Machine Learning": "https://www.coursera.org/learn/machine-learning",
    "Deep Learning": "https://www.deeplearning.ai/short-courses/",
    "Data Analysis": "https://www.datacamp.com/courses/data-analysis-in-python",
    "Statistics": "https://www.khanacademy.org/math/statistics-probability",
    "Data Visualization": "https://www.tableau.com/learn/training",
    "Pandas": "https://pandas.pydata.org/docs/getting_started/index.html",
    "NumPy": "https://numpy.org/devdocs/user/quickstart.html",
    "TensorFlow": "https://www.tensorflow.org/tutorials",
    "PyTorch": "https://pytorch.org/tutorials/",
    "scikit-learn": "https://scikit-learn.org/stable/tutorial/index.html",
    # Cloud & DevOps
    "AWS": "https://aws.amazon.com/training/",
    "Azure": "https://learn.microsoft.com/en-us/training/azure/",
    "Google Cloud": "https://cloud.google.com/training",
    "Docker": "https://docs.docker.com/get-started/",
    "Kubernetes": "https://kubernetes.io/docs/tutorials/",
    "API Development": "https://www.postman.com/resources/courses/",
    # Web/App
    "HTML": "https://developer.mozilla.org/en-US/docs/Web/HTML",
    "CSS": "https://developer.mozilla.org/en-US/docs/Web/CSS",
    "Sass": "https://sass-lang.com/guide",
    "React": "https://reactjs.org/tutorial/tutorial.html",
    "Angular": "https://angular.io/start",
    "Vue": "https://vuejs.org/guide/quick-start.html",
    "Next.js": "https://nextjs.org/learn/basics/create-nextjs-app",
    "Django": "https://docs.djangoproject.com/en/stable/intro/tutorial01/",
    "Flask": "https://flask.palletsprojects.com/en/2.0.x/tutorial/",
    "Firebase": "https://firebase.google.com/docs",
    "WordPress": "https://learn.wordpress.org/",
    # Collaboration
    "Git": "https://git-scm.com/doc",
    "GitHub": "https://docs.github.com/en/get-started",
    "JIRA": "https://www.atlassian.com/software/jira/guides",
    "Slack": "https://slack.com/resources",
    "Trello": "https://trello.com/guide",
    "Notion": "https://www.notion.so/help/guides/notion-basics",
    "Asana": "https://academy.asana.com/",
    # Productivity & Office
    "Excel": "https://support.microsoft.com/en-us/excel",
    "Google Sheets": "https://support.google.com/docs/answer/14078417",
    "PowerPoint": "https://support.microsoft.com/en-us/powerpoint",
    "Microsoft Word": "https://support.microsoft.com/en-us/word",
    # Content Creation Tools
    "Canva": "https://www.canva.com/learn/getting-started-with-canva/",
    "Figma": "https://www.figma.com/resources/learn-design/",
    "Photoshop": "https://helpx.adobe.com/photoshop/tutorials.html",
    "Illustrator": "https://helpx.adobe.com/illustrator/tutorials.html",
    "Premiere Pro": "https://helpx.adobe.com/premiere-pro/tutorials.html",
    "Final Cut Pro": "https://support.apple.com/final-cut-pro",
    "DaVinci Resolve": "https://www.blackmagicdesign.com/products/davinciresolve/training",
    "OBS Studio": "https://obsproject.com/wiki/",
    "Audacity": "https://www.audacityteam.org/help/documentation/",
    # AI & Generative
    "Prompt Engineering": "https://www.deeplearning.ai/short-courses/chatgpt-prompt-engineering-for-developers/",
    "Natural Language Processing": "https://www.coursera.org/specializations/natural-language-processing",
    "ChatGPT": "https://platform.openai.com/docs/guides/gpt",
    "Llama": "https://llama.meta.com/get-started/",
    "Generative AI": "https://developers.google.com/machine-learning/genai",
    "LangChain": "https://langchain.com/docs/",
    "Stable Diffusion": "https://docs.stablediffusionweb.com/",
    "Midjourney": "https://docs.midjourney.com/",
    "D-ID": "https://docs.d-id.com/docs",
    "Speech-to-Text": "https://cloud.google.com/speech-to-text/docs",
    "Text-to-Speech": "https://cloud.google.com/text-to-speech/docs",
    "Face Animation": "https://d-id.com/resources/",
    "Voice Cloning": "https://www.descript.com/overdub",
    "OpenCV": "https://opencv.org/courses/",
    "Text Summarization": "https://huggingface.co/tasks/summarization",
    "Topic Modeling": "https://towardsdatascience.com/topic-modeling-with-python-5eedb6e2a55",
    "Automated Captioning": "https://www.rev.com/blog/automated-captioning",
    # Content, Marketing, Publishing
    "Content Automation": "https://zapier.com/blog/automate-content-creation/",
    "Content Repurposing": "https://www.contenthacker.com/content-repurposing/",
    "Newsletter Automation": "https://www.mailerlite.com/blog/newsletter-automation-guide",
    "Blog Writing": "https://www.wix.com/blog/post/how-to-write-a-blog-post",
    "SEO": "https://moz.com/beginners-guide-to-seo",
    "Digital Marketing": "https://www.coursera.org/specializations/digital-marketing",
    "Branding": "https://www.coursera.org/learn/branding-the-creative/",
    "Copywriting": "https://www.copyblogger.com/copywriting-101/",
    "Social Media Management": "https://www.hubspot.com/social-media-management",
    "YouTube Studio": "https://support.google.com/youtube/answer/9314474",
    "Podcast Editing": "https://www.buzzsprout.com/blog/podcast-editing-software-tools",
    "Online Courses": "https://www.udemy.com/courses/teacher/",
    "Analytics": "https://support.google.com/analytics/answer/1008015",
    "Google Analytics": "https://analytics.google.com/analytics/academy/",
    "Monetization": "https://www.youtube.com/creators/monetize",
    "Affiliate Marketing": "https://www.hubspot.com/affiliate-marketing",
    "Email Marketing": "https://mailchimp.com/email-marketing/",
    "E-commerce Platforms": "https://www.shopify.com/learn",
    # Creative
    "Storytelling": "https://openai.com/research/storytelling",
    "Scriptwriting": "https://www.scriptreaderpro.com/complete-guide-to-scriptwriting/",
    "Photo Retouching": "https://www.adobe.com/creativecloud/photography/discover/photo-retouching.html",
    "Video Thumbnails": "https://www.youtube.com/official/blog/en-GB/products/yt-creator/video-thumbnails/",
    "Livestreaming": "https://creators.youtube.com/learn/live-streaming/",
    "Online Collaboration": "https://www.notion.so/help/guides/collaborate",
    "Campaign Optimization": "https://support.google.com/google-ads/answer/6167122"
}



def extract_text(file):
    text = ""
    if not file:
        return ""
    filepath = file.name
    if filepath.endswith('.pdf'):
        reader = PdfReader(filepath)
        for page in reader.pages:
            txt = page.extract_text()
            if txt:
                text += txt + "\n"
    elif filepath.endswith('.docx'):
        doc = Document(filepath)
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text.lower()

def suggest_roles(skill):
    roles_map = {
        "Python": ["Backend Developer", "Data Analyst"],
        "Machine Learning": ["ML Engineer", "Data Scientist"],
        "AWS": ["Cloud Engineer", "DevOps Engineer"],
        "Docker": ["DevOps Engineer"],
        "Flask": ["Backend Developer", "Web Developer"],
        "SQL": ["Database Admin", "Data Analyst"],
        "Kubernetes": ["Cloud Architect"],
        "Git": ["Software Engineer"]
    }
    return roles_map.get(skill, [])

def generate_pdf(your_skills, matching_skills, missing_skills, learning_resources, career_suggestions, match_score):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, "AI Resume & Job Match Report", 0, 1, 'C')
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, f"Match Score: {match_score}%", 0, 1, 'C')
    pdf.ln(5)
    def write_section(title, content_list):
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, title, ln=1)
        pdf.set_font("Arial", size=11)
        if content_list:
            for item in content_list:
                pdf.multi_cell(0, 8, "- " + str(item))
        else:
            pdf.cell(0, 8, "None", ln=1)
        pdf.ln(5)
    write_section("Your Skills:", your_skills)
    write_section("Matching Skills:", matching_skills)
    write_section("Missing Skills:", missing_skills)
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 10, "Learning Resources:", ln=1)
    pdf.set_font("Arial", size=11)
    if learning_resources:
        for skill, link in learning_resources.items():
            pdf.multi_cell(0, 8, f"{skill}: {link}")
    else:
        pdf.cell(0, 8, "None", ln=1)
    pdf.ln(5)
    write_section("Career Suggestions:", career_suggestions)
    path = "resume_job_match_report.pdf"
    pdf.output(path)
    return path

def summarize_job_description(job_desc):
    if not job_desc.strip():
        return job_desc
    API_URL = "https://api-inference.huggingface.co/models/facebook/bart-large-cnn"
    headers = {"Authorization": "Bearer -copy your hugging face key here-"}
    try:
        response = requests.post(API_URL, headers=headers, json={"inputs": job_desc})
        if response.status_code == 200:
            return response.json()[0]['summary_text']
        return job_desc
    except:
        return job_desc

def calculate_match_score(matching_skills, missing_skills):
    total_relevant = len(matching_skills) + len(missing_skills)
    if total_relevant == 0:
        return 0
    return int((len(matching_skills) / total_relevant) * 100)

def extract_keywords(text, n=10):
    words = re.findall(r'\b\w{4,}\b', text.lower())
    common_words = set(['with', 'this', 'that', 'have', 'will', 'your', 'team', 'work'])
    filtered = [word for word in words if word not in common_words and not word.isdigit()]
    return [word for word, count in Counter(filtered).most_common(n)]

def generate_cover_letter(job_title, company, matching_skills):
    if not job_title:
        job_title = "the position"
    if not company:
        company = "your company"
    skills_list = ", ".join(matching_skills[:5]) if matching_skills else "relevant skills"
    return f"""
Dear Hiring Manager,

I'm excited to apply for the {job_title} position at {company}. 
With my experience in {skills_list}, I believe I'd be a great fit for this role.

[Your personal introduction here]

Looking forward to the opportunity to discuss how my skills align with your needs.

Best regards,
[Your Name]
"""

def analyze_resume(resume_file, job_desc):
    if not resume_file or not job_desc.strip():
        return [], [], [], {}, [], None, 0, []
    resume_text = extract_text(resume_file)
    job_text = job_desc.lower()
    your_skills = [skill for skill in SKILLS_DB if skill.lower() in resume_text]
    matching_skills = [skill for skill in your_skills if skill.lower() in job_text]
    missing_skills = [skill for skill in SKILLS_DB if skill.lower() in job_text and skill not in your_skills]
    learning_resources = {skill: RESOURCES.get(skill, "No resource found.") for skill in missing_skills}
    career_suggestions = sorted(set(role for skill in your_skills for role in suggest_roles(skill)))
    match_score = calculate_match_score(matching_skills, missing_skills)
    job_keywords = extract_keywords(job_text)
    pdf_path = generate_pdf(your_skills, matching_skills, missing_skills, learning_resources, career_suggestions, match_score)
    return (your_skills, matching_skills, missing_skills, learning_resources, 
            career_suggestions, pdf_path, match_score, job_keywords)

def format_skills(skills):
    if not skills:
        return "<p>No skills found.</p>"
    return "<ul>" + "".join(f"<li>{skill}</li>" for skill in skills) + "</ul>"

def format_resources(resources):
    if not resources:
        return "<p>No learning resources available.</p>"
    return "<ul>" + "".join(f'<li><strong>{skill}:</strong> <a href="{link}" target="_blank">{link}</a></li>' 
                           for skill, link in resources.items()) + "</ul>"

def format_careers(careers):
    if not careers:
        return "<p>No career suggestions.</p>"
    return "<ul>" + "".join(f"<li>{career}</li>" for career in careers) + "</ul>"

def format_keywords(keywords):
    if not keywords:
        return "<p>No keywords extracted.</p>"
    return "<div class='keywords'>" + " ".join(f"<span class='keyword'>{kw}</span>" for kw in keywords) + "</div>"

